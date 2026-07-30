"""スキルシート（XLSX / テキスト PDF）から LLM 用の経験テキストを抽出する。"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import datetime
from typing import Literal

from app.models import TalentSkillSheet
from app.skill_sheet_names import extension_from_filename_or_mime

ExperienceExtractStatus = Literal["ok", "empty", "unsupported", "failed"]

MAX_EXPERIENCE_CHARS = 12_000
PROMPT_EXPERIENCE_CHARS = 6_000

_SECTION_HEADING_RE = re.compile(
    r"(経歴|職務経歴|プロジェクト|業務|参画|開発実績|work\s*history|project)",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class ExperienceExtractResult:
    status: ExperienceExtractStatus
    text: str | None = None


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _row_line(cells: list[str]) -> str:
    parts = [c for c in cells if c]
    if not parts:
        return ""
    return " | ".join(parts)


def _truncate(text: str, *, limit: int) -> str:
    cleaned = re.sub(r"\n{3,}", "\n\n", text.strip())
    if len(cleaned) <= limit:
        return cleaned
    return cleaned[: limit - 1] + "…"


def _extract_xlsx(data: bytes) -> ExperienceExtractResult:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ExperienceExtractResult(status="failed")

    try:
        wb = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        return ExperienceExtractResult(status="failed")

    sections: list[str] = []
    fallback_lines: list[str] = []

    for sheet in wb.worksheets:
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell_to_str(v) for v in row]
            if any(cells):
                rows.append(cells)

        if not rows:
            continue

        capture = False
        sheet_lines: list[str] = []
        for cells in rows:
            line = _row_line(cells)
            if not line:
                continue
            joined = " ".join(cells)
            if _SECTION_HEADING_RE.search(joined):
                capture = True
            if capture:
                sheet_lines.append(line)
            fallback_lines.append(line)

        if sheet_lines:
            header = f"--- {sheet.title} ---"
            sections.append(header + "\n" + "\n".join(sheet_lines))

    wb.close()

    if sections:
        text = _truncate("\n\n".join(sections), limit=MAX_EXPERIENCE_CHARS)
        return ExperienceExtractResult(status="ok", text=text or None)

    if fallback_lines:
        text = _truncate("\n".join(fallback_lines), limit=MAX_EXPERIENCE_CHARS)
        return ExperienceExtractResult(status="ok", text=text or None)

    return ExperienceExtractResult(status="empty")


def _extract_pdf(data: bytes) -> ExperienceExtractResult:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ExperienceExtractResult(status="failed")

    try:
        reader = PdfReader(io.BytesIO(data))
        parts: list[str] = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        text = "\n".join(parts).strip()
    except Exception:  # noqa: BLE001
        return ExperienceExtractResult(status="failed")

    if not text:
        return ExperienceExtractResult(status="empty")
    return ExperienceExtractResult(
        status="ok",
        text=_truncate(text, limit=MAX_EXPERIENCE_CHARS),
    )


def extract_experience_from_bytes(
    data: bytes,
    *,
    filename: str | None = None,
    content_type: str | None = None,
) -> ExperienceExtractResult:
    """ファイルバイト列から経験テキストを抽出する。"""
    ext = extension_from_filename_or_mime(filename, content_type)
    if ext in {"xlsx", "xlsm"}:
        return _extract_xlsx(data)
    if ext == "pdf":
        return _extract_pdf(data)
    if ext == "xls":
        return ExperienceExtractResult(status="unsupported")
    # content_type から推定
    mime = (content_type or "").lower()
    if "pdf" in mime:
        return _extract_pdf(data)
    if "spreadsheet" in mime or "excel" in mime or ext == "xlsx":
        return _extract_xlsx(data)
    return ExperienceExtractResult(status="unsupported")


def apply_experience_extract_to_row(row: TalentSkillSheet, result: ExperienceExtractResult) -> None:
    row.experience_text = result.text
    row.experience_extract_status = result.status
    row.experience_extracted_at = datetime.now().astimezone()


def extract_and_apply_to_row(row: TalentSkillSheet, data: bytes) -> ExperienceExtractResult:
    result = extract_experience_from_bytes(
        data,
        filename=row.filename,
        content_type=row.content_type,
    )
    apply_experience_extract_to_row(row, result)
    return result


def experience_text_for_prompt(text: str | None) -> str | None:
    if not text or not text.strip():
        return None
    return _truncate(text.strip(), limit=PROMPT_EXPERIENCE_CHARS)
